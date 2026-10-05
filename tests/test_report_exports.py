"""The client report PDF and the data workbook."""

import io

import pytest
from openpyxl import load_workbook
from sqlmodel import select

from api.config import get_settings
from api.models import Tenant
from tests.test_dashboards import member, processed_run  # noqa: F401  (fixtures)
from tests.test_processing import configured_tenant  # noqa: F401
from tests.test_runs import fake_retrieve, job_env, make_tenant  # noqa: F401  (fixtures)


@pytest.fixture()
def demo_tenant(db_session, monkeypatch, tmp_path):
    from api.demo_client import SLUG, build_demo_client

    monkeypatch.setattr(get_settings(), "raw_storage_dir", str(tmp_path))
    build_demo_client(db_session)
    return db_session.exec(select(Tenant).where(Tenant.slug == SLUG)).one()


def test_report_html_carries_every_section(db_session, demo_tenant):
    from api.report_pdf import render_report_html

    html = render_report_html(db_session, demo_tenant)
    for heading in ("Where Northpeak stands", "Each engine", "Why the numbers",
                    "What to do next", "Did the fixes work?"):
        assert heading in html
    assert "Fix a wrong fact: Starter plan" in html       # the playbook
    assert "/help/plans" in html                          # the fact conflict
    assert "<svg" in html                                 # the trend chart
    assert "Down 15.3 pts" in html and "pill-bad" in html  # bad news in ink


def test_report_pdf_renders(db_session, demo_tenant):
    from api.report_pdf import build_report_pdf

    pdf = build_report_pdf(db_session, demo_tenant)
    assert pdf.startswith(b"%PDF")
    assert len(pdf) > 20_000


def test_report_survives_an_empty_tenant(db_session):
    from api.report_pdf import build_report_pdf

    tenant = Tenant(name="Brand — “New” Co", slug="empty-co")
    db_session.add(tenant)
    db_session.commit()
    assert build_report_pdf(db_session, tenant).startswith(b"%PDF")


def test_falls_back_to_basic_pdf_without_renderer(db_session, monkeypatch):
    import api.report_pdf as report_pdf
    from api.reports import build_summary_pdf

    def broken(*_a, **_k):
        raise OSError("cannot load library 'libpango-1.0-0'")

    monkeypatch.setattr(report_pdf, "build_report_pdf", broken)
    tenant = Tenant(name="Fallback Co", slug="fallback-co")
    db_session.add(tenant)
    db_session.commit()
    assert build_summary_pdf(db_session, tenant).startswith(b"%PDF")


def test_workbook_tabs_and_formula_safety(db_session, demo_tenant):
    from api.report_xlsx import build_workbook

    wb = load_workbook(io.BytesIO(build_workbook(db_session, demo_tenant)))
    assert wb.sheetnames == [
        "Summary", "Engines", "Head to head", "Visibility by day", "Answers", "Cited sites",
        "Power pages", "Fan-out", "Agent picks", "Fact conflicts", "Playbook", "Proof",
    ]
    assert wb["Answers"].max_row > 500
    assert wb["Fact conflicts"]["C2"].value == "$19"
    assert isinstance(wb["Engines"]["B2"].value, float)


def test_workbook_never_writes_formulas(db_session):
    from openpyxl import Workbook

    from api.report_xlsx import _put

    ws = Workbook().active
    assert ws is not None
    _put(ws, 1, 1, '=HYPERLINK("http://evil","x")')
    assert ws["A1"].data_type == "s"


def test_workbook_route(client, login, member, processed_run):  # noqa: F811
    login(member, org_id="org_smith")
    resp = client.get("/api/tenant/reports/report.xlsx")
    assert resp.status_code == 200
    assert resp.content[:2] == b"PK"
    assert "spreadsheetml" in resp.headers["content-type"]
