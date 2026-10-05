"""The data workbook: one .xlsx with a tab per dataset, for analysts and
agency decks. Same sources as the PDF report and the dashboard."""

from __future__ import annotations

import csv
import io
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet
from sqlmodel import Session

from api.models import Tenant
from api.report_pdf import _section, report_context

HEAD_FILL = PatternFill("solid", fgColor="0B1220")
HEAD_FONT = Font(bold=True, color="FFFFFF")
TITLE_FONT = Font(bold=True, size=14)


def _put(ws: Worksheet, row: int, column: int, value: Any) -> None:
    cell = ws.cell(row=row, column=column, value=value)
    # openpyxl treats any string starting with "=" as a formula; data from
    # answers and pages must never execute in someone's spreadsheet.
    if isinstance(value, str) and value.startswith("="):
        cell.data_type = "s"


def _table(ws: Worksheet, headers: list[str], rows: list[list[Any]], start: int = 1,
           widths: dict[int, int] | None = None) -> None:
    for c, h in enumerate(headers, 1):
        cell = ws.cell(row=start, column=c, value=h)
        cell.fill, cell.font = HEAD_FILL, HEAD_FONT
    for r, row in enumerate(rows, start + 1):
        for c, v in enumerate(row, 1):
            _put(ws, r, c, v)
    if start == 1:
        ws.freeze_panes = "A2"
        if rows:
            ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(rows) + 1}"
    for c, h in enumerate(headers, 1):
        longest = max([len(str(h))] + [len(str(r[c - 1])) for r in rows[:200] if c - 1 < len(r)])
        width = (widths or {}).get(c, min(60, max(10, longest + 2)))
        ws.column_dimensions[get_column_letter(c)].width = width


def _csv_rows(text: str) -> tuple[list[str], list[list[Any]]]:
    rows = list(csv.reader(io.StringIO(text)))
    if not rows:
        return [], []

    def typed(v: str) -> Any:
        # CSV builders prefix risky text with ' ; restore it (cells are typed).
        if v.startswith("'") and v[1:2] in ("=", "+", "-", "@"):
            return v[1:]
        try:
            return int(v)
        except ValueError:
            try:
                return float(v)
            except ValueError:
                return {"True": True, "False": False}.get(v, v)

    return rows[0], [[typed(v) for v in r] for r in rows[1:]]


def build_workbook(session: Session, tenant: Tenant) -> bytes:
    from api.reports import build_results_csv, build_visibility_csv

    assert tenant.id is not None
    ctx = report_context(session, tenant)
    s, c, d = ctx["s"] or {}, ctx["c"] or {}, ctx["d"] or {}
    wb = Workbook()

    # Summary
    ws = wb.active
    assert ws is not None
    ws.title = "Summary"
    ws["A1"] = f"{ctx['brand']} · AI answer visibility"
    ws["A1"].font = TITLE_FONT
    ws["A2"] = f"{ctx['period']} · exported {ctx['generated']} · EverPresent"
    overall = s.get("overall") or {}
    verdict = s.get("verdict") or {}
    summary = [
        ["Verdict", verdict.get("line", "")],
        ["Named in answers (%)", overall.get("rate")],
        ["Range low (%)", overall.get("low")],
        ["Range high (%)", overall.get("high")],
        ["Answers measured", overall.get("answers")],
        ["Change vs previous two weeks", (overall.get("change") or {}).get("verdict", "")],
        ["Answer share (%)", s.get("answer_share")],
        ["Share of voice (%)", s.get("brand_sov")],
        ["Leads the answer (%)", (s.get("prominence") or {}).get("lead_rate")],
        ["Consistency", (s.get("stability") or {}).get("label", "")],
    ]
    _table(ws, ["Measure", "Value"], summary, start=4, widths={1: 32, 2: 90})
    ws.cell(row=len(summary) + 6, column=1, value="Share of voice").font = Font(bold=True)
    _table(ws, ["Brand", "Share (%)"], [[b["name"], b["value"]] for b in s.get("sov_bars") or []],
           start=len(summary) + 7, widths={1: 32, 2: 90})
    for row in ws.iter_rows(min_row=5, max_col=2):
        row[1].alignment = Alignment(wrap_text=True, vertical="top")

    # Engines
    ws = wb.create_sheet("Engines")
    _table(ws, ["Engine", "Named (%)", "Range low (%)", "Range high (%)", "Answers", "Mentioned",
                "Change", "Change (pts)", "p-value"],
           [[e.get("label"), e.get("rate"), e.get("low"), e.get("high"), e.get("answers"),
             e.get("mentioned"), (e.get("change") or {}).get("verdict"),
             (e.get("change") or {}).get("delta"), (e.get("change") or {}).get("p_value")]
            for e in s.get("engines") or []])

    # Head to head
    ws = wb.create_sheet("Head to head")
    _table(ws, ["Competitor", "Answers with both", f"{ctx['brand']} named first", "Win rate (%)"],
           [[h["competitor"], h["shared"], h["wins"], h["win_rate"]]
            for h in s.get("head_to_head") or []])

    # Visibility (daily) and Results (per answer)
    for title, text in (
        ("Visibility by day", _section("vis", build_visibility_csv, session, tenant.id)),
        ("Answers", _section("res", build_results_csv, session, tenant.id)),
    ):
        ws = wb.create_sheet(title)
        headers, rows = _csv_rows(text or "")
        _table(ws, headers, rows)

    # Citations
    full = _section("citations_full", _full_citations, session, tenant.id) or {}
    ws = wb.create_sheet("Cited sites")
    _table(ws, ["Site", "Type", "Citations", "Engines"],
           [[x["domain"], x.get("category"), x.get("count"), len(x.get("surfaces") or [])]
            for x in full.get("domains") or c.get("domains") or []])
    ws = wb.create_sheet("Power pages")
    _table(ws, ["Page", "Site", "Type", "Citations", "Questions", "Engines"],
           [[p["url"], p.get("domain"), p.get("category"), p.get("citations"), p.get("queries"),
             len(p.get("surfaces") or [])]
            for p in full.get("power_pages") or c.get("power_pages") or []])

    # Fan-out
    ws = wb.create_sheet("Fan-out")
    _table(ws, ["Question", "Sub-search", "Engines"],
           [[p["query"], q["text"], ", ".join(q.get("engines") or [])]
            for p in (full.get("fanout") or []) for q in p.get("subqueries") or []])

    # Agent picks
    picks = d.get("picks") or {}
    ws = wb.create_sheet("Agent picks")
    _table(ws, ["Prompt", "Type", "Picked first (%)", "Named (%)", "Answers", "Picked most"],
           [[p.get("text"), p.get("kind"), p["first_pick"]["rate"], p["named"]["rate"],
             p["named"]["answers"], p.get("leader")]
            for p in picks.get("prompts") or []])

    # Fact conflicts
    ws = wb.create_sheet("Fact conflicts")
    _table(ws, ["Page", "Fact", "Page says", "Fact sheet says", "Times cited", "Quote"],
           [[f["path"], f.get("subject"), f.get("stated"), f.get("label") or f.get("expected"),
             f.get("cited"), f.get("snippet")]
            for f in full.get("conflicts") or d.get("conflicts") or []])

    # Playbook
    ws = wb.create_sheet("Playbook")
    _table(ws, ["Priority", "Play", "Evidence", "Status", "What to do", "Steps", "Why"],
           [[r.priority, r.title, r.evidence, str(r.status), r.action_text,
             "\n".join(f"{i}. {st}" for i, st in enumerate(r.steps or [], 1)), r.why]
            for r in full.get("plays") or ctx["plays"]],
           widths={2: 40, 5: 60, 6: 80, 7: 60})
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    # Proof
    proof = ctx["proof"] or {}
    ws = wb.create_sheet("Proof")
    _table(ws, ["Fix", "Shipped", "Status", "Signal", "Before", "After", "Change", "Control change",
                "Compared with"],
           [[it.get("description"), it.get("shipped_at"), it.get("status"), g.get("label"),
             g.get("before"), g.get("after"),
             g.get("delta_pts") if g.get("kind") == "rate" else g.get("change_pct"),
             g.get("control_delta_pts") if g.get("kind") == "rate" else g.get("control_change_pct"),
             g.get("note")]
            for it in proof.get("items") or [] for g in it.get("signals") or []])

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def _full_citations(session: Session, tenant_id: int) -> dict:
    """The workbook carries everything, not the report's top-N cut."""
    from sqlmodel import col, select

    from api.dashboards_service import citations_intel, fanout_report
    from api.models import Recommendation, RecommendationStatus
    from api.own_pages_service import own_pages

    ci = citations_intel(session, tenant_id)
    pages = own_pages(session, tenant_id)
    plays = list(session.exec(
        select(Recommendation)
        .where(Recommendation.tenant_id == tenant_id,
               col(Recommendation.status).in_([RecommendationStatus.open,
                                               RecommendationStatus.in_progress]))
        .order_by(col(Recommendation.priority).desc())
    ).all())
    return {
        "domains": ci.get("domains") or [],
        "power_pages": ci.get("power_pages") or [],
        "fanout": fanout_report(session, tenant_id).get("prompts") or [],
        "conflicts": [{"path": p["path"], "cited": p.get("cited", 0), **f}
                      for p in pages.get("pages") or [] for f in p.get("fact_conflicts") or []],
        "plays": plays,
    }
